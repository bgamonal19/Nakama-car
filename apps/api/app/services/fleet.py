"""Customer fleet profile: every vehicle of a customer with its jobs to do and jobs done,
and the bulk import of a fleet from the Excel/CSV file the customer hands over."""
import csv
import io
import re
import unicodedata
from datetime import date, datetime

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.estimating import Estimate, EstimateStatus
from app.models.garage import Customer, RepairCase, RepairCaseStatus, Vehicle
from app.models.workshop import WorkOrder, WorkOrderTask, WorkTaskStatus
from app.services import tracking
from app.services.maintenance import maintenance_status

MAX_IMPORT_ROWS = 2000
MAX_IMPORT_BYTES = 5 * 1024 * 1024
DONE_LIMIT = 50

# Column -> accepted header names (compared without accents, case, spaces or punctuation).
COLUMNS: dict[str, tuple[str, ...]] = {
    "license_plate": ("targa", "matricula", "placa", "plate", "licenseplate", "patente"),
    "fleet_number": ("nflotta", "numeroflotta", "flotta", "ninterno", "interno", "nflota", "numeroflota", "flota", "fleetnumber", "fleet", "codice"),
    "make": ("marca", "make", "brand", "marcaveicolo", "costruttore", "casacostruttrice", "fabbricante", "produttore", "fabricante"),
    "model": ("modello", "modelo", "model", "modelloveicolo", "modeloveiculo", "modelovehiculo"),
    # "Fiat Ducato 2.3" in one cell: split into make and model when those columns are missing.
    "description": ("marcaemodello", "marcamodello", "marcaymodelo", "veicolo", "vehiculo", "descrizione", "descrizioneveicolo",
                    "descripcion", "mezzo", "automezzo", "denominazione", "makemodel", "vehicle", "description"),
    "version": ("versione", "version", "allestimento"),
    "year": ("anno", "ano", "year", "annoimmatricolazione", "immatricolazione"),
    "vin": ("telaio", "vin", "bastidor", "ntelaio", "chassis"),
    "mileage": ("km", "chilometri", "kilometri", "kilometros", "chilometraggio", "kilometraje", "mileage", "kmattuali"),
    "vehicle_category": ("tipo", "tipoveicolo", "tipovehiculo", "categoria", "category", "type"),
    "fuel_type": ("alimentazione", "carburante", "combustible", "fuel"),
    "color_name": ("colore", "color", "colour"),
    "engine_size": ("cilindrata", "cilindrada", "cc"),
    "power_kw": ("potenza", "potenzakw", "kw", "potencia", "potenciakw", "power"),
    "service_interval_km": ("tagliandoogni", "tagliandoognikm", "intervallokm", "ognikm", "intervalokm", "serviceintervalkm"),
    "service_interval_months": ("tagliandoognimesi", "intervallomesi", "ognimesi", "intervalomeses", "serviceintervalmonths"),
    "last_service_date": ("ultimotagliando", "dataultimotagliando", "ultimarevision", "ultimomantenimiento", "lastservice", "lastservicedate"),
    "last_service_km": ("kmultimotagliando", "kmultimarevision", "kmultimomantenimiento", "lastservicekm"),
}
TEMPLATE_HEADERS = [
    ("license_plate", "Targa"), ("fleet_number", "N. flotta"), ("make", "Marca"), ("model", "Modello"),
    ("version", "Versione"), ("year", "Anno"), ("vin", "Telaio"), ("mileage", "Km"),
    ("vehicle_category", "Tipo"), ("fuel_type", "Alimentazione"), ("color_name", "Colore"),
    ("service_interval_km", "Tagliando ogni km"), ("service_interval_months", "Tagliando ogni mesi"),
    ("last_service_date", "Ultimo tagliando"), ("last_service_km", "Km ultimo tagliando"),
]
TEMPLATE_EXAMPLE = ["AB123CD", "U-17", "Fiat", "Ducato", "2.3 MJT 140CV", 2021, "ZFA25000001234567", 85000, "Furgone",
                    "Diesel", "Bianco", 30000, 12, "2026-03-15", 80000]

CATEGORIES = {
    "CAR": ("car", "auto", "autovettura", "vettura", "coche", "turismo", "berlina", "suv"),
    "VAN": ("van", "furgone", "furgoneta", "furgon", "veicolocommerciale", "commerciale", "minibus"),
    "TRUCK": ("truck", "camion", "autocarro", "camión"),
    "TRACTOR": ("tractor", "trattore", "trattorestradale", "cabezatractora"),
    "TRAILER": ("trailer", "rimorchio", "semirimorchio", "remolque", "semirremolque"),
    "BUS": ("bus", "autobus", "pullman"),
    "MOTORCYCLE": ("motorcycle", "moto", "motociclo", "scooter", "ciclomotore"),
    "OTHER": ("other", "altro", "otro"),
}
# Headers that start with these words are read even when the rest differs ("Marca veicolo", "Modello / versione").
PREFIXES = (("targa", "license_plate"), ("matricola", "license_plate"), ("marca", "make"), ("modell", "model"), ("telaio", "vin"))
TEXT_LIMITS = {"description": 280, "fleet_number": 40, "vin": 32, "fuel_type": 40, "engine_size": 20, "make": 120, "model": 120, "version": 160, "color_name": 120}
INT_RANGES = {
    "year": (1886, 2100), "mileage": (0, 10_000_000), "power_kw": (0, 2000),
    "service_interval_km": (0, 500_000), "service_interval_months": (0, 120), "last_service_km": (0, 10_000_000),
}


class ImportFileError(ValueError):
    pass


def key(text: object) -> str:
    """Header/category text reduced to lowercase letters and digits, without accents."""
    value = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", value.lower())


def normalize_plate(value: object) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


# ---------- profile ----------

def vehicle_jobs(db: Session, vehicle: Vehicle) -> dict:
    """Jobs to do (open work, requests waiting for a work order, due maintenance) and jobs done."""
    tenant_id = vehicle.tenant_id
    cases = db.scalars(
        select(RepairCase)
        .where(RepairCase.vehicle_id == vehicle.id, RepairCase.tenant_id == tenant_id)
        .order_by(RepairCase.created_at.desc())
    ).all()
    rows = db.execute(
        select(WorkOrderTask, WorkOrder.repair_case_id)
        .join(WorkOrder, WorkOrder.id == WorkOrderTask.work_order_id)
        .where(WorkOrder.repair_case_id.in_([case.id for case in cases]), WorkOrder.tenant_id == tenant_id, WorkOrderTask.tenant_id == tenant_id)
        .order_by(WorkOrder.created_at, WorkOrderTask.sort_order, WorkOrderTask.created_at)
    ).all() if cases else []
    by_case = {case.id: case for case in cases}
    with_tasks = {case_id for _, case_id in rows}

    todo, done = [], []
    maintenance = maintenance_status(vehicle)
    if maintenance["state"] in ("DUE", "SOON"):
        todo.append({
            "kind": "MAINTENANCE", "description": "Tagliando", "status": maintenance["state"],
            "case_number": None, "case_id": None, "date": maintenance["next_service_date"], "km": maintenance["next_service_km"], "note": None,
        })
    for case in cases:
        if not tracking.is_closed(case) and case.id not in with_tasks:
            todo.append({
                "kind": "REQUEST", "description": case.customer_notes or "Pratica aperta", "status": case.status.value,
                "case_number": case.case_number, "case_id": str(case.id), "date": case.opened_at or case.created_at, "km": case.mileage, "note": None,
            })
    for task, case_id in rows:
        case = by_case[case_id]
        item = {
            "kind": "TASK", "description": task.description, "status": task.status.value,
            "case_number": case.case_number, "case_id": str(case.id), "km": case.mileage, "note": task.note,
        }
        if task.status == WorkTaskStatus.DONE:
            done.append({**item, "date": task.completed_at or task.updated_at})
        elif not tracking.is_closed(case):
            todo.append({**item, "date": case.opened_at or case.created_at})
    # Newest first in both lists.
    todo.sort(key=lambda item: str(item["date"] or ""), reverse=True)
    done.sort(key=lambda item: str(item["date"] or ""), reverse=True)
    # In the workshop: an open case not yet ready for pick-up (same rule as the portal's "In officina").
    in_workshop = any(not tracking.is_closed(case) and case.status != RepairCaseStatus.READY for case in cases)
    # Last time the vehicle came into the workshop (newest case), for sorting the fleet by entry.
    last_entry = max((case.opened_at or case.created_at for case in cases), default=None, key=lambda value: tracking.aware(value).timestamp())
    return {"todo": todo, "done": done[:DONE_LIMIT], "done_total": len(done), "in_workshop": in_workshop, "last_entry": last_entry}


def fleet_profile(db: Session, customer: Customer) -> dict:
    vehicles = db.scalars(
        select(Vehicle)
        .where(Vehicle.customer_id == customer.id, Vehicle.tenant_id == customer.tenant_id)
        .order_by(Vehicle.fleet_number, Vehicle.license_plate)
    ).all()
    items = []
    for vehicle in vehicles:
        jobs = vehicle_jobs(db, vehicle)
        items.append({
            "id": str(vehicle.id), "license_plate": vehicle.license_plate, "fleet_number": vehicle.fleet_number,
            "make": vehicle.make, "model": vehicle.model, "version": vehicle.version, "year": vehicle.year,
            "vin": vehicle.vin, "mileage": vehicle.mileage, "vehicle_category": vehicle.vehicle_category,
            "fuel_type": vehicle.fuel_type, "maintenance": maintenance_status(vehicle), **jobs,
        })
    return {
        "customer_id": str(customer.id),
        "customer_name": tracking.customer_display_name(customer),
        "totals": {
            "vehicles": len(items),
            "with_todo": sum(1 for item in items if item["todo"]),
            "todo": sum(len(item["todo"]) for item in items),
            "done": sum(item["done_total"] for item in items),
        },
        "vehicles": items,
    }


# Words that mark a case as routine maintenance rather than a repair (checked in requests and tasks).
MAINTENANCE_WORDS = (
    "tagliando", "manutenzione", "revisione", "cambioolio", "olio", "filtri", "filtro", "service",
    "mantenimiento", "revision", "aceite", "pneumatic", "gomme", "cambiogomme", "neumatic",
)


def case_kind(case: RepairCase, descriptions: list[str]) -> str:
    """MAINTENANCE for services (tagliando, oil, filters, tyres), otherwise REPAIR."""
    text = key(" ".join([case.customer_notes or "", *descriptions]))
    return "MAINTENANCE" if any(word in text for word in MAINTENANCE_WORDS) else "REPAIR"


def vehicle_profile(db: Session, vehicle: Vehicle) -> dict:
    """Card of one vehicle: its data, maintenance plan and every case with the work done, newest first."""
    tenant_id = vehicle.tenant_id
    customer = db.get(Customer, vehicle.customer_id) if vehicle.customer_id else None
    cases = db.scalars(
        select(RepairCase)
        .where(RepairCase.vehicle_id == vehicle.id, RepairCase.tenant_id == tenant_id)
        .order_by(RepairCase.created_at.desc())
    ).all()
    tasks: dict = {case.id: [] for case in cases}
    if cases:
        for task, case_id in db.execute(
            select(WorkOrderTask, WorkOrder.repair_case_id)
            .join(WorkOrder, WorkOrder.id == WorkOrderTask.work_order_id)
            .where(WorkOrder.repair_case_id.in_(list(tasks)), WorkOrder.tenant_id == tenant_id, WorkOrderTask.tenant_id == tenant_id)
            .order_by(WorkOrder.created_at, WorkOrderTask.sort_order, WorkOrderTask.created_at)
        ).all():
            tasks[case_id].append(task)
    totals: dict = {}
    if cases:
        for estimate in db.scalars(
            select(Estimate)
            .where(Estimate.repair_case_id.in_(list(tasks)), Estimate.tenant_id == tenant_id,
                   Estimate.status.not_in([EstimateStatus.DRAFT, EstimateStatus.SUPERSEDED]))
            .order_by(Estimate.created_at)
        ).all():
            totals[estimate.repair_case_id] = str(estimate.total)  # the latest one wins
    history = []
    for case in cases:
        items = tasks[case.id]
        history.append({
            "id": str(case.id), "case_number": case.case_number, "status": case.status.value,
            "work_type": case.work_type or "BODY", "kind": case_kind(case, [task.description for task in items]),
            "closed": tracking.is_closed(case), "opened_at": case.opened_at or case.created_at,
            "mileage": case.mileage, "customer_notes": case.customer_notes, "estimate_total": totals.get(case.id),
            "tasks": [
                {"description": task.description, "status": task.status.value, "note": task.note, "completed_at": task.completed_at}
                for task in items
            ],
        })
    return {
        "id": str(vehicle.id), "license_plate": vehicle.license_plate, "fleet_number": vehicle.fleet_number,
        "make": vehicle.make, "model": vehicle.model, "version": vehicle.version, "year": vehicle.year,
        "vin": vehicle.vin, "mileage": vehicle.mileage, "vehicle_category": vehicle.vehicle_category,
        "fuel_type": vehicle.fuel_type, "engine_size": vehicle.engine_size, "power_kw": vehicle.power_kw,
        "color_name": vehicle.color_name,
        "customer_id": str(customer.id) if customer else None,
        "customer_name": tracking.customer_display_name(customer) if customer else None,
        "maintenance": maintenance_status(vehicle),
        "totals": {
            "cases": len(history),
            "repairs": sum(1 for item in history if item["kind"] == "REPAIR"),
            "maintenance": sum(1 for item in history if item["kind"] == "MAINTENANCE"),
            "open": sum(1 for item in history if not item["closed"]),
        },
        "history": history,
    }


# ---------- import ----------

def read_table(content: bytes, filename: str) -> list[list[object]]:
    if not content:
        raise ImportFileError("Il file è vuoto.")
    if len(content) > MAX_IMPORT_BYTES:
        raise ImportFileError("Il file è troppo grande (massimo 5 MB).")
    name = (filename or "").lower()
    if name.endswith(".xls"):
        raise ImportFileError("Formato .xls non supportato: salva il file come .xlsx o .csv.")
    if name.endswith((".xlsx", ".xlsm")) or content[:2] == b"PK":
        from openpyxl import load_workbook
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # corrupted or not really an Excel file
            raise ImportFileError("Impossibile leggere il file Excel.") from exc
        sheet = workbook.worksheets[0]
        table = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        return table
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def find_header(table: list[list[object]]) -> tuple[int, dict[int, str]]:
    lookup = {alias: field for field, aliases in COLUMNS.items() for alias in aliases}
    for index, row in enumerate(table[:15]):
        mapping: dict[int, str] = {}
        for column, cell in enumerate(row):
            header = key(cell)
            field = lookup.get(header) or next((name for prefix, name in PREFIXES if header.startswith(prefix)), None)
            if field and field not in mapping.values():
                mapping[column] = field
        if "license_plate" in mapping.values():
            return index, mapping
    raise ImportFileError("Non trovo la colonna «Targa» nelle prime righe del file.")


def category(value: object) -> str | None:
    text = key(value)
    if not text:
        return None
    for code, aliases in CATEGORIES.items():
        if text == code.lower() or text in aliases:
            return code
    return None


def parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise ValueError


def parse_int(value: object) -> int:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(value)
    text = re.sub(r"[\s.']", "", str(value)).replace(",", ".")
    text = re.sub(r"(?i)(km|kw|cc)$", "", text)
    return int(float(text))


def blank(row: list[object]) -> bool:
    return all(cell is None or not str(cell).strip() for cell in row)


def clean_row(row: list[object], mapping: dict[int, str]) -> tuple[dict, list[str]]:
    data: dict = {}
    warnings: list[str] = []
    for column, field in mapping.items():
        raw = row[column] if column < len(row) else None
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            continue
        if field == "license_plate":
            data[field] = normalize_plate(raw)
        elif field == "vehicle_category":
            code = category(raw)
            if code:
                data[field] = code
            else:
                warnings.append(f"tipo «{raw}» non riconosciuto")
        elif field == "last_service_date":
            try:
                data[field] = parse_date(raw)
            except ValueError:
                warnings.append(f"data «{raw}» non valida")
        elif field in INT_RANGES:
            low, high = INT_RANGES[field]
            try:
                number = parse_int(raw)
                if not low <= number <= high:
                    raise ValueError
                data[field] = number
            except ValueError:
                warnings.append(f"{field} «{raw}» non valido")
        else:
            text = str(int(raw)) if isinstance(raw, float) and raw.is_integer() else str(raw).strip()
            if field == "vin":
                text = text.upper().replace(" ", "")
            data[field] = text[:TEXT_LIMITS.get(field, 255)]
    return data, warnings


def split_description(data: dict) -> None:
    """Fill make and model from a single "vehicle" text when the file has no separate columns."""
    words = (data.pop("description", "") or "").split()
    if not words:
        return
    if not data.get("make"):
        data["make"] = words.pop(0)[:TEXT_LIMITS["make"]]
    elif words and key(words[0]) == key(data["make"]):
        words.pop(0)
    if words and not data.get("model"):
        data["model"] = " ".join(words)[:TEXT_LIMITS["model"]]


def import_fleet(db: Session, customer: Customer, content: bytes, filename: str, dry_run: bool) -> dict:
    """Create or update the customer's vehicles from the file; with dry_run nothing is saved."""
    table = read_table(content, filename)
    header, mapping = find_header(table)
    body = [(number, row) for number, row in enumerate(table[header + 1:], start=header + 2) if not blank(row)]
    if not body:
        raise ImportFileError("Il file non contiene veicoli.")
    if len(body) > MAX_IMPORT_ROWS:
        raise ImportFileError(f"Troppe righe: massimo {MAX_IMPORT_ROWS} veicoli per file.")

    existing = {
        normalize_plate(vehicle.license_plate): vehicle
        for vehicle in db.scalars(select(Vehicle).where(Vehicle.tenant_id == customer.tenant_id)).all()
    }
    owners: dict = {}
    seen: set[str] = set()
    results = []
    counts = {"created": 0, "updated": 0, "skipped": 0}
    for number, row in body:
        data, warnings = clean_row(row, mapping)
        split_description(data)
        plate = data.pop("license_plate", "")
        result = {"row": number, "license_plate": plate, "warnings": warnings}
        # A plate always has a digit: rows like "AZIENDA" or "TOTALE" are titles or notes, not vehicles.
        if not 2 <= len(plate) <= 20 or not any(char.isdigit() for char in plate):
            results.append({**result, "action": "skipped", "message": "Targa mancante o non valida"})
            counts["skipped"] += 1
            continue
        if plate in seen:
            results.append({**result, "action": "skipped", "message": "Targa ripetuta nel file"})
            counts["skipped"] += 1
            continue
        seen.add(plate)
        vehicle = existing.get(plate)
        if vehicle is not None and vehicle.customer_id not in (None, customer.id):
            if vehicle.customer_id not in owners:
                owners[vehicle.customer_id] = db.get(Customer, vehicle.customer_id)
            owner = tracking.customer_display_name(owners[vehicle.customer_id])
            results.append({**result, "action": "skipped", "message": f"Targa già registrata per il cliente {owner}"})
            counts["skipped"] += 1
            continue
        if vehicle is None:
            vehicle = Vehicle(tenant_id=customer.tenant_id, customer_id=customer.id, license_plate=plate, vehicle_category=data.pop("vehicle_category", "CAR"))
            db.add(vehicle)
            existing[plate] = vehicle
            action = "created"
        else:
            vehicle.customer_id = customer.id
            action = "updated"
        for field, value in data.items():
            setattr(vehicle, field, value)
        results.append({**result, "action": action, "message": None})
        counts[action] += 1

    if dry_run:
        db.rollback()
    else:
        db.commit()
    return {
        "dry_run": dry_run,
        "columns": [field for field in mapping.values()],
        **counts,
        "rows": results,
    }


def template_workbook() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Flotta"
    sheet.append([label for _, label in TEMPLATE_HEADERS])
    sheet.append(TEMPLATE_EXAMPLE)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = max(12, len(str(column[0].value)) + 4)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
