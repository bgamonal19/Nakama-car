import io
from datetime import date, timedelta

from openpyxl import Workbook, load_workbook

import test_fleet_billing as fleet
import test_tracking as tracking_tests
from test_users import client  # noqa: F401  (pytest fixture)


def excel(rows):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_fleet_profile_lists_jobs_to_do_and_done_by_plate(client):
    c, _, _ = client
    customer, _, vehicle, case = fleet.create_fleet_case(c)
    order = tracking_tests.open_work_order(c, case)
    tasks = c.get(f"/api/v1/work-orders/{order['id']}/tasks").json()
    assert c.patch(f"/api/v1/work-orders/{order['id']}/tasks/{tasks[0]['id']}/status", json={"status": "DONE"}).status_code == 200
    waiting = c.post("/api/v1/vehicles", json={"license_plate": "GH456JK", "customer_id": customer["id"], "vehicle_category": "VAN"}).json()
    c.post("/api/v1/cases", json={"customer_id": customer["id"], "vehicle_id": waiting["id"], "customer_notes": "Spia motore accesa"})
    c.patch(f"/api/v1/vehicles/{waiting['id']}/maintenance", json={
        "mileage": 61000, "service_interval_km": 30000, "last_service_km": 30000,
        "last_service_date": (date.today() - timedelta(days=30)).isoformat(),
    })

    profile = c.get(f"/api/v1/customers/{customer['id']}/fleet")
    assert profile.status_code == 200, profile.text
    data = profile.json()
    assert data["customer_name"] == "Univex Group" and data["totals"]["vehicles"] == 2
    by_plate = {item["license_plate"]: item for item in data["vehicles"]}
    truck = by_plate[vehicle["license_plate"]]
    assert [job["description"] for job in truck["done"]] == [tasks[0]["description"]]
    assert all(job["kind"] == "TASK" and job["case_number"] == case["case_number"] for job in truck["todo"])
    assert len(truck["todo"]) == len(tasks) - 1
    van = by_plate["GH456JK"]
    assert {job["kind"] for job in van["todo"]} == {"REQUEST", "MAINTENANCE"}
    assert any(job["description"] == "Spia motore accesa" for job in van["todo"])
    # Both have an open case, so both are in the workshop (red row in the list).
    assert truck["in_workshop"] and van["in_workshop"]
    assert truck["last_entry"] and van["last_entry"]

    # The customer sees the same profile from the portal.
    account = c.post(f"/api/v1/customers/{customer['id']}/portal-accounts", json={"email": "flotta@univex.example.com", "full_name": "Ufficio flotta"}).json()
    portal = c.__class__(c.app)
    token = portal.post("/api/v1/portal/auth/login", json={"email": account["email"], "password": account["temporary_password"]}).json()["access_token"]
    portal.headers["Authorization"] = "Bearer " + token
    assert portal.get("/api/v1/portal/fleet").json()["totals"] == data["totals"]


def test_fleet_import_from_excel_preview_then_save(client):
    c, _, _ = client
    customer, _, vehicle, _ = fleet.create_fleet_case(c)
    other = c.post("/api/v1/customers", json={"first_name": "Mario", "last_name": "Rossi"}).json()
    c.post("/api/v1/vehicles", json={"license_plate": "ZZ999ZZ", "customer_id": other["id"]})
    content = excel([
        ["Elenco flotta Univex"],
        ["Targa", "N° flotta", "Marca", "Modello", "Anno", "Km", "Tipo", "Ultimo tagliando", "Tagliando ogni km"],
        ["ab 123 cd", "U-01", "Fiat", "Ducato", 2021, "85.000", "Furgone", "15/03/2026", 30000],
        [vehicle["license_plate"], "U-17", "Iveco", "Daily", 2019, 120000, "Camion", None, None],
        ["ZZ999ZZ", "U-02", "Ford", "Transit", 2020, 1000, "Van", None, None],
        ["AB123CD", "U-03", "Fiat", "Doblò", 2022, 1000, "Auto", None, None],
        [None, "U-04", "Fiat", "Panda", "nuova", None, "Astronave", None, None],
        [],
    ])
    url = f"/api/v1/customers/{customer['id']}/fleet/import?filename=flotta.xlsx"
    preview = c.post(url + "&dry_run=true", content=content, headers={"Content-Type": "application/octet-stream"})
    assert preview.status_code == 200, preview.text
    result = preview.json()
    assert (result["created"], result["updated"], result["skipped"]) == (1, 1, 3)
    assert [row["action"] for row in result["rows"]] == ["created", "updated", "skipped", "skipped", "skipped"]
    assert "Mario Rossi" in result["rows"][2]["message"]
    assert result["rows"][3]["message"] == "Targa ripetuta nel file"
    assert c.get(f"/api/v1/customers/{customer['id']}/fleet").json()["totals"]["vehicles"] == 1  # nothing saved yet

    saved = c.post(url, content=content, headers={"Content-Type": "application/octet-stream"}).json()
    assert saved["created"] == 1 and saved["dry_run"] is False
    vehicles = {v["license_plate"]: v for v in c.get(f"/api/v1/customers/{customer['id']}/fleet").json()["vehicles"]}
    assert set(vehicles) == {"AB123CD", vehicle["license_plate"]}
    new = vehicles["AB123CD"]
    assert (new["make"], new["mileage"], new["vehicle_category"], new["fleet_number"]) == ("Fiat", 85000, "VAN", "U-01")
    assert new["maintenance"]["last_service_date"] == "2026-03-15" and new["maintenance"]["next_service_km"] is None
    assert vehicles[vehicle["license_plate"]]["make"] == "Iveco" and vehicles[vehicle["license_plate"]]["vehicle_category"] == "TRUCK"


def test_fleet_import_csv_template_and_errors(client):
    c, _, _ = client
    customer, _, _, _ = fleet.create_fleet_case(c)
    url = f"/api/v1/customers/{customer['id']}/fleet/import"
    csv_file = "Matrícula;Marca;Modelo;Kilómetros\nXY987WV;Renault;Master;12 500\n".encode("latin-1")
    result = c.post(url + "?filename=flota.csv", content=csv_file, headers={"Content-Type": "text/csv"}).json()
    assert result["created"] == 1 and result["rows"][0]["license_plate"] == "XY987WV"
    assert c.post(url + "?filename=flota.xls", content=b"data").status_code == 422
    missing = c.post(url + "?filename=flota.csv", content=b"Nome;Marca\nA;B\n")
    assert missing.status_code == 422 and "Targa" in missing.json()["detail"]

    template = c.get("/api/v1/fleet-import/template")
    assert template.status_code == 200
    headers = [cell.value for cell in load_workbook(io.BytesIO(template.content)).active[1]]
    assert headers[0] == "Targa"
    reimport = c.post(url + "?filename=modello.xlsx&dry_run=true", content=template.content).json()
    assert reimport["created"] == 1 and not reimport["rows"][0]["warnings"]


def test_maintenance_is_soon_within_5000_km():
    from app.models.garage import Vehicle
    from app.services.maintenance import maintenance_status

    def state(mileage):
        vehicle = Vehicle(mileage=mileage, last_service_km=100000, service_interval_km=20000)
        return maintenance_status(vehicle)["state"]

    assert state(114999) == "OK"
    assert state(115000) == "SOON"
    assert state(120000) == "DUE"
